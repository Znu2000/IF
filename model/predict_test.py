"""
[PRED] 과제1 테스트A 예측 → 제출 파일
  1) 학습 선박 전체로 핵심 앙상블(lgbm_base + catboost + hier_hyp)을 다시 학습
     피처 = 전체 fold 랭킹(feature_ranking_{pool}.csv) 상위 k*  (CV 와 같은 선정 규칙)
  2) 클래스 가중 보정(+adj): 저장된 OOF 전체로 가중을 한 번 맞춰서 테스트 확률에 곱함
     리더보드 지표가 micro-F1(= 정확도)이라 기본은 정확도 기준 (--metric macro 로 바꿀 수 있음)
  3) 테스트 선박 피처(test_ship_features.parquet, 2025년 항적만으로 계산)로 예측
     ※ 학습 데이터와 같은 rfid_id 가 있어도 ID 로 정답을 찾아오지 않음 (항적만으로 판별)
  4) 제출샘플.xlsx 를 복사해 테스트A 시트를 채움 → results/submission_testA.xlsx
     테스트B 시트는 과제2 전이라 rfid_id 만 남기고 비움
  5) 분포 이동 점검: 학습(2026) vs 테스트(2025) 선박 피처를 구분하는 분류기 AUC (0.5 에 가까울수록 같음)

선행: python fe_hourly.py --src "<테스트셋 폴더>" --prefix test_hourly
      python fe_aggregate.py --test
실행: python predict_test.py [--pool B]
"""
import argparse
import shutil

import lightgbm as lgb
import numpy as np
import openpyxl
import pandas as pd
from sklearn.metrics import accuracy_score, roc_auc_score
from sklearn.model_selection import cross_val_predict

import cv_common as V
import fe_config as C
from train_classifiers import MODELS, fit_class_weights

CORE = ["lgbm_base", "catboost", "hier_hyp"]
SAMPLE = C.PROJECT_ROOT / "문제 데이터_수정_1027" / "테스트셋 및 결과제출 샘플" / "제출샘플.xlsx"


def select_feats(pool, k):
    r = pd.read_csv(C.RESULT_DIR / f"feature_ranking_{pool}.csv")
    r = r[r["redundant_folds"] <= 2].sort_values("final_rank")  # 다수 fold 에서 중복 제거된 피처는 제외
    return r["feature"].head(k).tolist()


def shift_check(S, T, cols):
    X = pd.concat([S[cols], T[cols]], ignore_index=True)
    d = np.r_[np.zeros(len(S)), np.ones(len(T))]
    m = lgb.LGBMClassifier(n_estimators=200, learning_rate=0.05, num_leaves=31, max_bin=63, verbose=-1,
                           n_jobs=C.N_JOBS)
    p = cross_val_predict(m, X, d, cv=5, method="predict_proba")[:, 1]
    imp = pd.Series(m.fit(X, d).feature_importances_, index=cols).sort_values(ascending=False)
    return roc_auc_score(d, p), imp.head(10)


def main(a):
    S, y = V.load_ship()
    T = pd.read_parquet(C.FEAT_DIR / "test_ship_features.parquet")
    R = __import__("json").loads((C.RESULT_DIR / "fold_rankings.json").read_text(encoding="utf-8"))
    k = R[a.pool]["k_star"]
    feats = select_feats(a.pool, k)
    pool_cols = V.feature_pool(S, a.pool)
    hyp = [c for c in feats + pool_cols if V.is_hyp(c)]
    hyp = list(dict.fromkeys(hyp))[:30]
    print(f"Pool {a.pool}: k*={k} 피처, 가설 피처 {len(hyp)}개 / 학습 {len(S):,}척, 테스트 {len(T):,}척")

    Xtr, Xte = S[feats].reset_index(drop=True), T[feats].reset_index(drop=True)
    Xtr.attrs["full"], Xte.attrs["full"] = S[pool_cols].reset_index(drop=True), T[pool_cols].reset_index(drop=True)
    P = np.mean([MODELS[m](Xtr, y, Xte, hyp_cols=hyp) for m in CORE], axis=0)

    oof = np.mean([np.load(C.RESULT_DIR / f"oof_{a.pool}_{m}.npy") for m in CORE], axis=0)
    score = accuracy_score if a.metric == "micro" else V.macro_f1
    w = fit_class_weights(oof, y, score=score)
    print("클래스 가중:", {c: round(float(v), 3) for c, v in zip(C.CLASSES, w)},
          f"(OOF {a.metric}-F1 {score(y, oof.argmax(1)):.4f} → {score(y, (oof * w).argmax(1)):.4f})")
    pred = (P * w).argmax(1)
    out = pd.DataFrame({"rfid_id": T["rfid_id"].astype(str),
                        "pred": np.array(C.CLASSES)[pred],
                        "prob": P[np.arange(len(P)), pred]})  # 보정 전 앙상블 확률 (예측 클래스의 확률)
    for i, c in enumerate(C.CLASSES):
        out[f"p_{c}"] = P[:, i]
    out.round(4).to_csv(C.RESULT_DIR / "test_pred_all.csv", index=False, encoding="utf-8-sig")

    prior = pd.Series(y).map(dict(enumerate(C.CLASSES))).value_counts(normalize=True)
    print("\n예측 분포 (테스트 전체) vs 학습 비율\n",
          pd.DataFrame({"test_pred": out["pred"].value_counts(normalize=True), "train": prior}).round(3))

    # 제출 파일
    dst = C.RESULT_DIR / "submission_testA.xlsx"
    shutil.copy(SAMPLE, dst)
    wb = openpyxl.load_workbook(dst)
    wa, wbB = wb.worksheets[0], wb.worksheets[1]
    look = out.set_index("rfid_id")
    fallback = prior.idxmax()
    miss = []
    for row in wa.iter_rows(min_row=2):
        rid = str(row[0].value).strip()
        if rid in look.index:
            row[1].value, row[2].value = look.at[rid, "pred"], round(float(look.at[rid, "prob"]), 4)
        else:  # 테스트 항적이 없거나 파싱 불가 → 학습 최빈 클래스
            row[1].value, row[2].value = fallback, round(float(prior[fallback]), 4)
            miss.append(rid)
    for row in wbB.iter_rows(min_row=2):
        for cell in row[1:]:
            cell.value = None
    wb.save(dst)
    ids = [str(r[0].value).strip() for r in wa.iter_rows(min_row=2)]
    sub = look.reindex(ids)
    print(f"\n테스트A {len(ids)}척 → {dst.name}  (피처 없음 {len(miss)}척: {miss})")
    print(sub["pred"].value_counts().to_string())
    print(f"예측 확률 중앙값 {sub['prob'].median():.3f}")

    auc, top = shift_check(S, T, feats)
    print(f"\n분포 이동 점검 (학습 2026 vs 테스트 2025 구분 AUC): {auc:.3f}")
    print("구분에 많이 쓰인 피처:", ", ".join(top.index))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--pool", default="B", choices=["A", "B"])
    ap.add_argument("--metric", default="micro", choices=["micro", "macro"], help="리더보드 = micro-F1")
    main(ap.parse_args())
