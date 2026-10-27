"""
[DOC] 팀 공유용 피처 카탈로그 엑셀 → results/SeD_피처목록.xlsx
  시트: 읽어보기 / 범주요약 / 범주구조 / 선박피처 / PoolB_최종30 / 일단위_MIL / 시간단위_L1
  분류·설명은 feature_taxonomy.py, 순위는 results/feature_ranking_{A,B}.csv
실행: python make_feature_catalog.py
"""
import json
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import cv_common as V
import fe_config as C
import feature_taxonomy as T
from fe_hourly import FEAT_COLS
from mil import DAY_DROP, DAY_LOC

ARIAL = "Arial"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
HL_FILL = PatternFill("solid", fgColor="FFF2CC")
HYP_NAME = {"A": "재방문", "B": "직선 투망", "C": "속도 구조", "D": "선회", "E": "하루 리듬", "G": "이질성", "H": "계절"}


def table(ws, r0, header, rows, widths=None, wrap_cols=()):
    """r0 행부터 표 쓰기 → 다음 빈 행 번호"""
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.border = Font(name=ARIAL, bold=True, color="FFFFFF"), HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(i, j, v)
            c.font, c.border = Font(name=ARIAL, size=10), BORDER
            c.alignment = Alignment(vertical="top", wrap_text=(j in wrap_cols))
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return r0 + len(rows) + 2


def section(ws, r, title):
    c = ws.cell(r, 1, title)
    c.font, c.fill = Font(name=ARIAL, bold=True, size=12, color="1F4E78"), SEC_FILL
    for j in range(2, 7):
        ws.cell(r, j).fill = SEC_FILL
    return r + 1


def text(ws, r, s, bold=False):
    c = ws.cell(r, 1, s)
    c.font = Font(name=ARIAL, size=10, bold=bold)
    return r + 1


def data_sheet(wb, title, df, widths, wrap_cols=(), hl_col=None):
    ws = wb.create_sheet(title)
    table(ws, 1, list(df.columns), list(df.itertuples(index=False)), widths, wrap_cols)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    if hl_col:
        k = list(df.columns).index(hl_col) + 1
        for row in ws.iter_rows(min_row=2):
            if row[k - 1].value == "O":
                for cell in row:
                    cell.fill = HL_FILL
    return ws


def hyp_tag(c):
    """가설 피처면 해당 가설 그룹 표시"""
    if not V.is_hyp(c):
        return ""
    b = T.base_of(c)[0]
    src = T.describe_base(b)[1]
    for k in HYP_NAME:
        if f"가설 {k}" in src:
            return f"가설 {k}({HYP_NAME[k]})"
    cat = T.BASE_TO_TAX[b][0]
    for k in HYP_NAME:
        if f"({k})" in cat:
            return f"가설 {k}({HYP_NAME[k]})"
    return "가설 G(이질성: 날짜 간 표준편차)" if c.endswith("_std") else "가설 관련"


def main():
    S, _ = V.load_ship()
    cols = [c for c in S.columns if c not in V.KEYS]
    R = json.loads((C.RESULT_DIR / "fold_rankings.json").read_text(encoding="utf-8"))
    kA, kB = R["A"]["k_star"], R["B"]["k_star"]
    rk = {p: pd.read_csv(C.RESULT_DIR / f"feature_ranking_{p}.csv").set_index("feature") for p in "AB"}
    topB = list(rk["B"].index[:kB])

    # ── 선박 피처 표
    rows = []
    for i, c in enumerate(cols, 1):
        b, how, unit = T.base_of(c)
        cat, sub = T.BASE_TO_TAX[b]
        desc, src = T.describe_base(b)
        a, bb = rk["A"].reindex([c]).iloc[0], rk["B"].reindex([c]).iloc[0]
        rows.append({
            "No": i, "피처명": c, "대분류": cat, "세부 구분": sub, "기본 피처": b, "기본 피처 설명": desc,
            "만드는 방법": how, "계산 단위": unit, "출처": src,
            "위치(LOC)": "O" if c in V.LOC_COLS else "", "가설": hyp_tag(c),
            "Pool A 후보": "" if c in V.LOC_COLS else "O", "Pool B 후보": "O",
            f"Pool B 최종{kB}": "O" if c in topB else "",
            "A 선정빈도": a.get("topk_freq"), "B 선정빈도": bb.get("topk_freq"),
            "A 순위": a.get("final_rank"), "B 순위": bb.get("final_rank"),
            "중복 시 대표(B)": bb.get("redundant_of") if isinstance(bb.get("redundant_of"), str) else "",
        })
    F = pd.DataFrame(rows)

    # ── 일 단위 / 시간 단위
    drows = []
    for c in [c for c in pd.read_parquet(C.FEAT_DIR / "shipday.parquet").columns
              if c not in DAY_DROP and c != "part_dt"]:
        if c in T.D:
            b, (d, src) = c, T.D[c]
        else:
            b = c[3:]
            d, src = T.H[b]
            d = "그날 이동 시간대 평균: " + d
        cat, sub = T.BASE_TO_TAX[b]
        drows.append({"피처명": c, "대분류": cat, "세부 구분": sub, "설명": d, "출처": src,
                      "위치(LOC)": "O" if c in DAY_LOC else "", "MIL Pool A": "" if c in DAY_LOC else "O",
                      "MIL Pool B": "O"})
    DF = pd.DataFrame(drows)
    HF = pd.DataFrame([{"피처명": c, "대분류": T.BASE_TO_TAX[c][0], "세부 구분": T.BASE_TO_TAX[c][1],
                        "설명": T.H[c][0], "출처": T.H[c][1]} for c in FEAT_COLS])

    wb = Workbook()
    # ══════════ 1. 읽어보기 ══════════
    ws = wb.active
    ws.title = "읽어보기"
    ws.column_dimensions["A"].width = 22
    for j, w in zip("BCDEF", [34, 14, 44, 20, 14]):
        ws.column_dimensions[j].width = w
    r = 1
    c = ws.cell(r, 1, "SeD팀 과제 1 (조업형태 분류) 피처 카탈로그")
    c.font = Font(name=ARIAL, bold=True, size=16, color="1F4E78")
    r = text(ws, r + 1, f"작성 {date.today():%Y-%m-%d} · 코드 model/feature_taxonomy.py, make_feature_catalog.py · "
                        "모든 수치는 선박 단위 5-fold 교차검증(rfid_id 기준 분리) 결과")
    r = text(ws, r, "팀원 4명(지유·민성·나연·진우)의 피처 명세와 진행1 §6 가설(A~H)을 하나의 파이프라인(model/)으로 다시 구현한 피처 전체 목록입니다.")
    r += 1

    r = section(ws, r, "1. 왜 피처가 346개인가: 데이터 흐름")
    r = table(ws, r, ["단계", "1행의 단위", "피처 수", "어떻게 만드나", "코드", "시트"], [
        ["L1 시간", "선박·날짜·시간대", len(FEAT_COLS), "1시간 항적(MultiLineString) 하나에서 직접 계산. 튐 제거·평활 등 정제 후", "fe_hourly.py", "시간단위_L1"],
        ["L2 선박-일", "선박·날짜", len(DF), "하루의 시간대들을 요약 + '이동 시간대'만 평균낸 mv_* 44개", "fe_aggregate.py", "일단위_MIL"],
        ["L3 선박", "선박", len(F), "출어일 값들의 평균·중앙값·표준편차(+일부 90% 분위) + 조업 시간대 풀링 + 선박 전용 → 분류 모델 입력", "fe_aggregate.py", "선박피처"],
    ], wrap_cols=(4,))
    r = text(ws, r - 1, f"→ 서로 다른 '기본 피처'는 {F['기본 피처'].nunique()}종뿐. 같은 기본 피처에 통계(평균/중앙값/표준편차…)를 여러 개 붙이면서 {len(F)}개가 됨.")
    r = text(ws, r, "→ 같은 기본 피처에서 나온 '형제 피처'는 선박피처 시트의 '기본 피처' 열로 필터해서 함께 볼 수 있습니다.")
    r += 1

    r = section(ws, r, "2. Pool A / Pool B: 위치 피처를 넣느냐 빼느냐")
    r = table(ws, r, ["구분", "포함 피처", "후보 수", "목적", "선정 결과(k*)", "최고 macro-F1"], [
        ["Pool A", "위치 7개 제외", int((F["Pool A 후보"] == "O").sum()), "순수 조업 '행동'만으로 판별 → 처음 보는 해역의 선박에도 통하는 모델",
         f"중복 제거 후 전부 ({kA}개)", "0.523 (앙상블+보정)"],
        ["Pool B", "위치 포함", int((F["Pool B 후보"] == "O").sum()), "대회 점수 극대화 (어느 바다에서 조업하는지도 활용)",
         f"상위 {kB}개", "0.540 (앙상블)"],
    ], wrap_cols=(4,))
    r = text(ws, r - 1, "왜 나누나: 위치를 넣으면 F1이 오르지만, 모델이 '조업 방식'이 아니라 '어느 지역 배인가'를 외우는 shortcut 위험이 있음 (HANDOFF §3.3). 그래서 두 모델을 항상 같이 보고.")
    r = text(ws, r, "위치 피처 7개: " + ", ".join(V.LOC_COLS) + "  (선박피처 시트 '위치(LOC)' = O)")
    r = text(ws, r, "Pool A가 피처를 줄이지 않은 이유: 위치를 빼면 신호가 약한 피처 여러 개에 얇게 퍼져 있어, 줄일수록 1 표준오차 넘게 성능이 떨어짐 (k=60 0.447 < 전체 0.461).")
    r += 1

    r = section(ws, r, "3. 분류 체계: 대분류 12 → 세부 구분 → 기본 피처 → 선박 피처")
    r = text(ws, r, f"대분류 12개, 세부 구분 {len(T.TAXONOMY)}개. 범주는 사람이 읽기 위한 이름표이고 모델에는 쓰이지 않음 (피처 하나하나를 따로 순위 매김).")
    r = text(ws, r, "출처: 팀원 명세의 그룹(민성 공간/방향/속도/시간, 나연 v2 사전 그룹, 지유 ①~⑦) + 진행1 §6 가설 그룹을 합쳐 정리. 전체 구조는 '범주구조' 시트.")
    r = table(ws, r, ["가설 표시", "뜻 (진행1 §6 자망 vs 연안복합)", "", "", "", ""], [
        ["A", "재방문: 자망은 어구를 놓고 떠났다가 같은 자리로 돌아와 회수"], ["B", "직선 투망: 그물을 수백 m~수 km 직선으로 펼침"],
        ["C", "속도 구조: 양망 시 긴 극저속, 업종별 속도 조합"], ["D", "선회: 루프·톱니형 재이동"],
        ["E", "하루 리듬: 투망-휴식-양망 블록 구조, 조업 시각"], ["G", "선박 내 이질성: 날마다 다른 어구 사용 (날짜 간 표준편차, 출어일 유형)"],
        ["H", "계절: 월별 조업 패턴"], ["(F)", "지역 → 위치(LOC) 피처로 별도 관리"],
    ])
    r = section(ws, r, "4. 피처 이름 읽는 법")
    r = table(ws, r, ["표기", "뜻", "예시", "", "", ""], [
        ["mv_XXX_mean / _median / _std", "출어일마다 '이동 시간대'의 XXX 평균을 구한 뒤 → 출어일들의 평균 / 중앙값 / 표준편차(날마다 얼마나 다른가)", "mv_turn_mean_median"],
        ["XXX_mean / _median / _std", "하루 단위 값 XXX 의 출어일 평균 / 중앙값 / 표준편차", "longest_streak_mean"],
        ["XXX_p90", "출어일 값의 90% 분위 (많이 나간 날 기준)", "dist_day_km_p90"],
        ["fh_XXX_q50 / _q90", "조업 시간대(이동 중 & 항해 아님) 1시간 값들의 선박별 중앙값 / 90% 분위 (민성 C)", "fh_pca_elong_q90"],
        ["ks_", "5개 점 간격(k-step)으로 잰 지터 강건 버전 (민성)", "ks_speed_med"],
        ["_sea", "항구 정박 시간을 뺀 나머지 시간 기준 비율 (민성 H4)", "drift_share_sea"],
        ["_share / _ratio / _frac", "비율 (0~1)", "port_share"],
        ["_m / _km / _kn / _h", "단위: 미터 / 킬로미터 / 노트 / 시간", "home_km_max"],
    ], wrap_cols=(2,))
    r = section(ws, r, "5. 피처 선정 방식 (모두 fold 안의 학습 데이터로만 계산 → 검증 선박 정보 누수 없음)")
    r = table(ws, r, ["단계", "무엇을", "왜", "", "", ""], [
        ["① 사전 필터", "결측 60% 초과·상수 피처 제거", "쓸 수 없는 피처 정리 (라벨 미사용)"],
        ["② 중요도 3종", "SHAP 40% + 순열 중요도 40% + 상호정보량(MI) 20% 의 순위 평균",
         "SHAP=모델 기여, 순열=빼면 얼마나 떨어지나, MI=모델과 무관한 단독 신호. 한 지표의 편향 상쇄"],
        ["③ 노이즈 검정", "값을 무작위로 섞은 가짜 피처와 비교 (Boruta 방식)", "팀 규칙 '노이즈 기준선보다 나은가'를 피처마다 확인"],
        ["④ 중복 제거", "순위 순서대로 Spearman |ρ|≥0.95 인 쌍은 상위 하나만", "같은 정보 반복 방지 → '중복 시 대표(B)' 열"],
        ["⑤ k 결정", "k 5~150·전체 스윕, 최고점 − 1 표준오차 안의 가장 작은 k (1-SE rule)", "차이가 통계적으로 없으면 단순한 쪽 선택"],
    ], wrap_cols=(2, 3))
    r = text(ws, r - 1, "선정빈도 = 5개 fold 중 그 fold 의 상위 k* 안에 든 비율. fold마다 학습 데이터로 따로 뽑아서 구성이 조금씩 다름 → 빈도가 높을수록 안정적인 피처.")
    r += 1
    r = section(ws, r, "6. 시트 안내")
    r = table(ws, r, ["시트", "내용", "", "", "", ""], [
        ["범주요약", "대분류별 피처 수 (전체 / Pool A 후보 / Pool B 후보 / B 최종)"],
        ["범주구조", "대분류 → 세부 구분 → 무엇을 재나 → 기본 피처 → 해당 선박 피처 목록 (세부 구분별 1행)"],
        ["선박피처", f"분류 모델 입력 피처 {len(F)}개 전체. 노란 행 = Pool B 최종 {kB}개. 필터로 범주·기본 피처별 보기"],
        ["PoolB_최종30", f"Pool B 에서 최종 선정된 {kB}개만 순위 순으로 한 줄 설명과 함께"],
        ["일단위_MIL", f"하루 단위 피처 {len(DF)}개 (MIL 모델 입력, Pool A 는 위치 8개 제외 {int((DF['MIL Pool A'] == 'O').sum())}개)"],
        ["시간단위_L1", f"1시간 원천 피처 {len(HF)}개와 출처"],
    ], wrap_cols=(2,))
    r = section(ws, r, "7. 선박피처 시트 열 설명")
    table(ws, r, ["열", "설명", "", "", "", ""], [
        ["대분류 / 세부 구분", "분류체계 (3번 참고)"], ["기본 피처", "이 피처의 원천이 되는 시간·일·선박 단위 값"],
        ["만드는 방법", "기본 피처에서 선박 피처로 바꾸는 집계 방식"], ["계산 단위", "L1→L2→L3 = 시간→하루→선박 순으로 집계"],
        ["출처", "명세를 낸 팀원 또는 가설 (여러 명이 같은 피처를 냈으면 함께 표기, 공통 = 전원)"],
        ["가설", "진행1 §6 가설 그룹 (3번 표)"], ["Pool A/B 후보", "각 풀의 후보에 들어가는지"],
        [f"Pool B 최종{kB}", "Pool B 모델에 최종 사용된 상위 30개 (전 fold 종합 순위 기준)"],
        ["A/B 선정빈도", f"5개 fold 중 상위 k*(A {kA}, B {kB}) 안에 든 비율"], ["A/B 순위", "전 fold 종합 순위 (1이 최상위)"],
        ["중복 시 대표(B)", "Spearman |ρ|≥0.95 로 이 피처 대신 쓰인 피처 (빈칸 = 중복 아님)"],
    ], wrap_cols=(2,))

    # ══════════ 2. 범주요약 (수식) ══════════
    n = len(F) + 1
    L = {h: get_column_letter(list(F.columns).index(h) + 1) for h in F.columns}
    rng = lambda h: f"선박피처!${L[h]}$2:${L[h]}${n}"
    w2 = wb.create_sheet("범주요약")
    cats = list(dict.fromkeys(cat for cat, *_ in T.TAXONOMY))
    rows2 = []
    for i, cat in enumerate(cats, 2):
        rows2.append([cat, sum(1 for c_, *_ in T.TAXONOMY if c_ == cat), f"=COUNTIF({rng('대분류')},A{i})",
                      f'=COUNTIFS({rng("대분류")},A{i},{rng("Pool A 후보")},"O")',
                      f'=COUNTIFS({rng("대분류")},A{i},{rng("Pool B 후보")},"O")',
                      f'=COUNTIFS({rng("대분류")},A{i},{rng(f"Pool B 최종{kB}")},"O")'])
    last = len(cats) + 1
    rows2.append(["합계"] + [f"=SUM({col}2:{col}{last})" for col in "BCDEF"])
    table(w2, 1, ["대분류", "세부 구분 수", "선박 피처 수", "Pool A 후보", "Pool B 후보", f"Pool B 최종{kB}"], rows2,
          [24, 12, 13, 13, 13, 14])
    for cell in w2[last + 1]:
        cell.font = Font(name=ARIAL, bold=True)
    w2.cell(last + 3, 1, "※ 숫자는 '선박피처' 시트를 COUNTIF 로 세는 수식 (선박피처를 수정하면 자동 갱신)").font = Font(name=ARIAL, size=9, italic=True)

    # ══════════ 3. 범주구조 ══════════
    w3 = wb.create_sheet("범주구조")
    rows3 = []
    for i, (cat, sub, what, bases) in enumerate(T.TAXONOMY, 2):
        sub_f = F[F["세부 구분"] == sub]
        srcs = sorted({s for b in bases for s in T.describe_base(b)[1].split("·")})
        rows3.append([cat, sub, what, ", ".join(srcs), ", ".join(bases),
                      f'=COUNTIFS({rng("대분류")},A{i},{rng("세부 구분")},B{i})',
                      f'=COUNTIFS({rng("대분류")},A{i},{rng("세부 구분")},B{i},{rng(f"Pool B 최종{kB}")},"O")',
                      ", ".join(sub_f["피처명"]),
                      ", ".join(sub_f.loc[sub_f[f"Pool B 최종{kB}"] == "O", "피처명"])])
    table(w3, 1, ["대분류", "세부 구분", "무엇을 재나", "출처", "기본 피처 (원천 값)", "선박 피처 수",
                  f"B 최종{kB} 포함", "선박 피처 목록", f"B 최종{kB} 포함 피처"], rows3,
          [20, 20, 40, 18, 46, 10, 10, 70, 40], wrap_cols=(3, 5, 8, 9))
    w3.freeze_panes = "C2"
    # 같은 대분류는 색 띠로 구분
    band = [PatternFill("solid", fgColor="F2F2F2"), PatternFill(fill_type=None)]
    k = 0
    for i in range(2, len(rows3) + 2):
        if i > 2 and w3.cell(i, 1).value != w3.cell(i - 1, 1).value:
            k ^= 1
        for j in range(1, 10):
            if band[k].fill_type:
                w3.cell(i, j).fill = band[k]

    # ══════════ 4~7. 데이터 시트 ══════════
    data_sheet(wb, "선박피처", F, [6, 30, 18, 18, 22, 40, 34, 11, 14, 8, 18, 8, 8, 9, 8, 8, 7, 7, 26],
               wrap_cols=(6, 7), hl_col=f"Pool B 최종{kB}")
    ws4 = wb["선박피처"]
    for col in ("A 선정빈도", "B 선정빈도"):
        j = list(F.columns).index(col) + 1
        for row in ws4.iter_rows(min_row=2, min_col=j, max_col=j):
            row[0].number_format = "0%"
    TB = F[F[f"Pool B 최종{kB}"] == "O"].sort_values("B 순위")
    TB = pd.DataFrame({"B 순위": TB["B 순위"], "피처명": TB["피처명"], "대분류": TB["대분류"], "세부 구분": TB["세부 구분"],
                       "한 줄 설명": TB["기본 피처 설명"] + "  ← " + TB["만드는 방법"], "출처": TB["출처"],
                       "위치(LOC)": TB["위치(LOC)"], "가설": TB["가설"], "5fold 선정빈도": TB["B 선정빈도"]})
    w5 = data_sheet(wb, "PoolB_최종30", TB, [7, 30, 18, 18, 70, 14, 8, 18, 10], wrap_cols=(5,))
    for row in w5.iter_rows(min_row=2, min_col=9, max_col=9):
        row[0].number_format = "0%"
    data_sheet(wb, "일단위_MIL", DF, [26, 18, 18, 56, 14, 8, 10, 10], wrap_cols=(4,))
    data_sheet(wb, "시간단위_L1", HF, [22, 18, 18, 60, 16], wrap_cols=(4,))

    wb.calculation.fullCalcOnLoad = True  # 수식은 엑셀이 열 때 계산
    out = C.RESULT_DIR / "SeD_피처목록.xlsx"
    wb.save(out)
    print(f"✅ {out}  (선박 {len(F)} / 기본 피처 {F['기본 피처'].nunique()} / 세부 구분 {len(T.TAXONOMY)} / 일 {len(DF)} / 시간 {len(HF)})")
    # 수식 검산용 기대값
    chk = F.groupby("대분류").agg(n=("피처명", "size"), a=("Pool A 후보", lambda s: (s == "O").sum()),
                               b=("Pool B 후보", lambda s: (s == "O").sum()),
                               top=(f"Pool B 최종{kB}", lambda s: (s == "O").sum())).reindex(cats)
    print(chk.to_string(), "\n합계", chk.sum().to_dict())
    print("세부 구분별 피처 수 합:", F.groupby(["대분류", "세부 구분"]).size().sum())


if __name__ == "__main__":
    main()
